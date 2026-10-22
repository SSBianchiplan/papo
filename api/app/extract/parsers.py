"""Extração de texto e tabelas. Funções puras (bytes → resultado), testáveis sem banco."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

CHUNK = 2800                 # caracteres por seção (Word/texto)
MAX_TEXT = 2_000_000         # caracteres guardados por documento
MAX_ROWS = 200_000           # linhas guardadas por aba
MAX_PAGES = 2000
MAX_COLS = 500


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[Any]]
    total: int


@dataclass
class Extracted:
    kind: str
    parts: list[tuple[str, str]] = field(default_factory=list)
    sheets: list[Sheet] = field(default_factory=list)
    pages: int | None = None
    note: str | None = None

    @property
    def readable(self) -> bool:
        return any(t.strip() for _, t in self.parts) or any(s.rows for s in self.sheets)


DOC_LIKE = re.compile(r"\.(pptx|ppsx|potx|odt|odp|rtf|epub|zip|skill)$")
TEXT_LIKE = re.compile(r"\.(eml|ipynb)$")


def kind_of(filename: str, mime: str) -> str:
    n = filename.lower()
    if DOC_LIKE.search(n):
        return "docx"            # documento em seções (slides, capítulos, arquivos do zip)
    if n.endswith(".ods"):
        return "table"
    if TEXT_LIKE.search(n):
        return "text"
    if n.endswith(".pdf") or mime == "application/pdf":
        return "pdf"
    if n.endswith(".docx"):
        return "docx"
    if re.search(r"\.(xlsx|xlsm|csv|tsv)$", n):
        return "table"
    if mime.startswith("text/") or re.search(
        r"\.(txt|md|markdown|json|xml|html?|log|ya?ml|sql|py|js|ts|tsx|jsx|css|ini|cfg|conf|sh|java|c|cpp|cs|go|rs|rb|php)$", n
    ):
        return "text"
    return "other"


def _chunk(text: str, size: int = CHUNK) -> list[tuple[str, str]]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    parts: list[str] = []
    cur = ""
    for p in re.split(r"\n{2,}", text):
        if len(p) > size:
            if cur:
                parts.append(cur)
                cur = ""
            parts.extend(p[i:i + size] for i in range(0, len(p), size))
            continue
        if cur and len(cur) + 2 + len(p) > size:
            parts.append(cur)
            cur = ""
        cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        parts.append(cur)
    return [(f"seção {i + 1}", t) for i, t in enumerate(parts) if t.strip()]


def _cap(parts: list[tuple[str, str]]) -> tuple[list[tuple[str, str]], bool]:
    out: list[tuple[str, str]] = []
    total = 0
    for label, t in parts:
        if total + len(t) > MAX_TEXT:
            return out, True
        out.append((label, t))
        total += len(t)
    return out, False


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat(timespec="minutes")
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, int | float):
        return v
    return str(v).strip()


def _sheet_from_matrix(name: str, matrix: list[list[Any]]) -> Sheet | None:
    matrix = [r for r in matrix if any(c not in ("", None) for c in r)]
    if not matrix:
        return None
    width = min(max(len(r) for r in matrix), MAX_COLS)
    # cabeçalho: primeira linha com pelo menos metade das colunas preenchidas
    hi = next((i for i, r in enumerate(matrix[:30]) if sum(1 for c in r if c not in ("", None)) >= max(1, (width + 1) // 2)), 0)
    header = [str(c).strip() if c not in ("", None) else f"Coluna {i + 1}" for i, c in enumerate((matrix[hi] + [""] * width)[:width])]
    seen: dict[str, int] = {}
    for i, h in enumerate(header):                       # nomes de coluna únicos
        k = h.lower()
        if k in seen:
            seen[k] += 1
            header[i] = f"{h} ({seen[k]})"
        else:
            seen[k] = 1
    body = matrix[hi + 1:]
    rows = [[_cell(c) for c in (r + [""] * width)[:width]] for r in body[:MAX_ROWS]]
    return Sheet(name=name[:120] or "Planilha", columns=header, rows=rows, total=len(body))


def parse_pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:  # noqa: BLE001 — pypdf lança tipos variados para senha
            raise ValueError("PDF protegido por senha") from e
    n = len(reader.pages)
    parts: list[tuple[str, str]] = []
    for i, page in enumerate(reader.pages[:MAX_PAGES], start=1):
        t = page.extract_text() or ""
        parts.append((f"p. {i}", re.sub(r"[ \t]+", " ", t).strip()))
    parts, cut = _cap(parts)
    out = Extracted(kind="pdf", parts=parts, pages=n)
    chars = sum(len(t) for _, t in parts)
    if chars < 8 * max(1, len(parts)):
        out.note = "PDF com pouco ou nenhum texto (provavelmente escaneado). Envie as páginas como imagem."
    elif cut or n > MAX_PAGES:
        out.note = f"Guardadas as primeiras {len(parts)} páginas."
    return out


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_mark(p: Any, txt: str) -> str:
    """Mantém a estrutura do Word no texto: título vira '## ', item de lista vira '• '."""
    ppr = p.find(f"{_W}pPr")
    style = ""
    if ppr is not None:
        st = ppr.find(f"{_W}pStyle")
        style = (st.get(f"{_W}val") or "") if st is not None else ""
        if ppr.find(f"{_W}numPr") is not None or re.match(r"(?i)list|lista|pargrafodalista", style):
            return f"• {txt.strip()}"
    m = re.match(r"(?i)^(heading|t[ií]tulo|title)\s*(\d?)$", style)
    if m:
        return f"{'#' * min(int(m.group(2) or 1) + 1, 4)} {txt.strip()}"
    return txt


def parse_docx(data: bytes) -> Extracted:
    import docx

    d = docx.Document(io.BytesIO(data))
    blocks: list[str] = []
    body = d.element.body
    # mantém a ordem real de parágrafos e tabelas
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            txt = "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t"))
            if txt.strip():
                blocks.append(_docx_mark(el, txt))
        elif tag == "tbl":
            lines = []
            for tr in el.iter():
                if tr.tag.endswith("}tr"):
                    cells = []
                    for tc in tr.iterchildren():
                        if tc.tag.endswith("}tc"):
                            cells.append(" ".join("".join(t.text or "" for t in p.iter() if t.tag.endswith("}t"))
                                                  for p in tc.iter() if p.tag.endswith("}p")).strip())
                    lines.append(" | ".join(cells))
            if lines:
                blocks.append("\n".join(lines))
    text = ""
    for i, b in enumerate(blocks):            # itens de lista seguidos ficam no mesmo bloco
        sep = "" if i == 0 else ("\n" if b.startswith("• ") and blocks[i - 1].startswith("• ") else "\n\n")
        text += sep + b
    parts, cut = _cap(_chunk(text))
    return Extracted(kind="docx", parts=parts, note="Documento longo: guardado o início." if cut else None)


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        return ";" if sample.count(";") > sample.count(",") else ","


def parse_table(data: bytes, filename: str) -> Extracted:
    n = filename.lower()
    sheets: list[Sheet] = []
    if n.endswith((".csv", ".tsv")):
        text = decode_text(data)
        delim = "\t" if n.endswith(".tsv") else _sniff_delimiter(text[:20000])
        matrix: list[list[Any]] = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
        s = _sheet_from_matrix(re.sub(r"\.[a-z]+$", "", filename) or "Planilha", matrix)
        if s:
            sheets.append(s)
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                matrix = []
                for r in ws.iter_rows(values_only=True):
                    matrix.append(list(r))
                    if len(matrix) > MAX_ROWS + 30:
                        break
                total_hint = ws.max_row or len(matrix)
                s = _sheet_from_matrix(ws.title, matrix)
                if s:
                    s.total = max(s.total, total_hint - 1 if total_hint else s.total)
                    sheets.append(s)
        finally:
            wb.close()
    out = Extracted(kind="table", sheets=sheets)
    if any(len(s.rows) < s.total for s in sheets):
        out.note = f"Planilha grande: guardadas até {MAX_ROWS:,} linhas por aba.".replace(",", ".")
    return out


def _sections(kind: str, parts: list[tuple[str, str]], pages: int | None = None) -> Extracted:
    parts, cut = _cap([(lab, txt) for lab, txt in parts if txt.strip()])
    return Extracted(kind=kind, parts=parts, pages=pages, note="Arquivo longo: guardado o início." if cut else None)


def _text(kind: str, txt: str) -> Extracted:
    return _sections(kind, _chunk(txt))


def parse_other_formats(data: bytes, filename: str) -> Extracted | None:
    """PowerPoint, OpenDocument, RTF, e-mail, Jupyter, EPUB, HTML e .zip. None = não é um destes."""
    from . import office

    n = filename.lower()
    parts: list[tuple[str, str]]
    if n.endswith((".pptx", ".ppsx", ".potx")):
        slides = office.pptx_slides(data)
        return _sections("docx", slides, pages=len(slides))
    if n.endswith(".odp"):
        slides = office.odp_slides(data)
        return _sections("docx", slides, pages=len(slides))
    if n.endswith(".odt"):
        return _text("docx", office.odt_text(data))
    if n.endswith(".ods"):
        sheets = [sh for sh in (_sheet_from_matrix(name, m) for name, m in office.ods_tables(data)) if sh]
        return Extracted(kind="table", sheets=sheets, note=None if sheets else "Planilha sem dados.")
    if n.endswith(".rtf"):
        return _text("docx", office.rtf_text(data))
    if n.endswith(".eml"):
        return _text("text", office.eml_text(data))
    if n.endswith(".ipynb"):
        return _text("text", office.ipynb_text(data))
    if n.endswith(".epub"):
        parts = []
        for chapter, txt in office.epub_chapters(data):
            parts.extend(_label(chapter, _chunk(txt)))
        return _sections("docx", parts)
    if n.endswith((".html", ".htm", ".xhtml")):
        return _text("text", office.html_to_text(decode_text(data)))
    if n.endswith((".zip", ".skill")):        # .skill do hub é lido como pacote: SKILL.md e referências
        parts = []
        skipped: list[str] = []
        for inner, blob in office.zip_members(data):
            try:
                ex = parse(blob, inner, "application/octet-stream")
            except Exception:  # noqa: BLE001 — um arquivo ruim dentro do zip não derruba os outros
                ex = Extracted(kind="other")
            got = [(f"{inner} · {lab}"[:60], txt) for lab, txt in ex.parts]
            for sh in ex.sheets:          # planilha dentro do zip entra como texto tabulado
                got.append((f"{inner} · {sh.name}"[:60], "\t".join(sh.columns) + "\n" + "\n".join("\t".join(str(c) for c in r) for r in sh.rows[:2000])))
            if got:
                parts.extend(got)
            else:
                skipped.append(inner)
        ex = _sections("docx", parts)
        if skipped:
            ex.note = ((ex.note + " ") if ex.note else "") + f"{len(skipped)} arquivo(s) do zip não lidos."
        return ex
    return None


def _label(prefix: str, parts: list[tuple[str, str]]) -> list[tuple[str, str]]:
    return [(prefix if len(parts) == 1 else f"{prefix} · {lab}", body) for lab, body in parts]


def parse(data: bytes, filename: str, mime: str) -> Extracted:
    other = parse_other_formats(data, filename)
    if other is not None:
        return other
    k = kind_of(filename, mime)
    if k == "pdf":
        return parse_pdf(data)
    if k == "docx":
        return parse_docx(data)
    if k == "table":
        if filename.lower().endswith(".xls"):
            return Extracted(kind="table", note="Formato .xls antigo não é lido. Salve como .xlsx ou CSV.")
        return parse_table(data, filename)
    if k == "text":
        parts, cut = _cap(_chunk(decode_text(data)))
        return Extracted(kind="text", parts=parts, note="Arquivo longo: guardado o início." if cut else None)
    low = filename.lower()
    if low.endswith(".doc"):
        return Extracted(kind="other", note="Formato .doc antigo não é lido. Salve como .docx ou PDF.")
    if low.endswith((".ppt", ".pps")):
        return Extracted(kind="other", note="Formato .ppt antigo não é lido. Salve como .pptx ou PDF.")
    if low.endswith((".heic", ".heif")):
        return Extracted(kind="other", note="Foto HEIC do iPhone não é lida. Envie como JPG (Ajustes → Câmera → Mais compatível).")
    if low.endswith(".xls"):
        return Extracted(kind="other", note="Formato .xls antigo não é lido. Salve como .xlsx ou CSV.")
    return Extracted(kind="other", note="Formato não lido pelo agente.")
