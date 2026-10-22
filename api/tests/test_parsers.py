"""Leitura de arquivos reais gerados no teste."""
from __future__ import annotations

import io

from app.extract.parsers import parse


def _pdf(pages: list[str]) -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for t in pages:
        c.drawString(72, 750, t)
        c.showPage()
    c.save()
    return buf.getvalue()


def test_pdf_por_pagina() -> None:
    ex = parse(_pdf(["Objeto do contrato: fornecimento de paineis eletricos para a obra.",
                     "O prazo de entrega e de 90 dias apos a assinatura do pedido."]), "contrato.pdf", "application/pdf")
    assert ex.kind == "pdf" and ex.pages == 2
    assert [label for label, _ in ex.parts] == ["p. 1", "p. 2"]
    assert "90 dias" in ex.parts[1][1] and ex.note is None


def test_pdf_sem_texto_sinalizado() -> None:
    ex = parse(_pdf(["", ""]), "scan.pdf", "application/pdf")
    assert ex.note and "escaneado" in ex.note and not ex.readable


def test_docx_paragrafos_e_tabela_em_ordem() -> None:
    import docx

    d = docx.Document()
    d.add_paragraph("Política de garantia")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Prazo"
    t.cell(1, 0).text, t.cell(1, 1).text = "Painel", "12 meses"
    d.add_paragraph("Fim do documento")
    buf = io.BytesIO()
    d.save(buf)
    ex = parse(buf.getvalue(), "garantia.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    text = "\n".join(t for _, t in ex.parts)
    assert text.index("Política") < text.index("Painel | 12 meses") < text.index("Fim do documento")


def test_xlsx_cabecalho_e_tipos() -> None:
    from datetime import date

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["Relatório mensal"])          # título acima do cabeçalho
    ws.append(["Região", "Valor", "Data"])
    ws.append(["Sul", 1200.5, date(2026, 1, 3)])
    ws.append(["Norte", 300, date(2026, 1, 4)])
    buf = io.BytesIO()
    wb.save(buf)
    ex = parse(buf.getvalue(), "vendas.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    s = ex.sheets[0]
    assert s.name == "Vendas" and s.columns == ["Região", "Valor", "Data"]
    assert s.rows == [["Sul", 1200.5, "2026-01-03"], ["Norte", 300, "2026-01-04"]]


def test_csv_brasileiro_ponto_e_virgula_latin1() -> None:
    raw = "Região;Valor\nSul;1.200,50\nSão Paulo;300\n".encode("cp1252")
    ex = parse(raw, "vendas.csv", "text/csv")
    s = ex.sheets[0]
    assert s.columns == ["Região", "Valor"] and s.rows[1] == ["São Paulo", "300"]


def test_formatos_antigos_explicados() -> None:
    assert "docx" in (parse(b"x", "a.doc", "application/msword").note or "")
    assert "xlsx" in (parse(b"x", "a.xls", "application/vnd.ms-excel").note or "")


# ---------------------------------------------------------------- formatos adicionais
def _zip(files: dict[str, str | bytes]) -> bytes:
    import io
    import zipfile

    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        for p, c in files.items():
            z.writestr(p, c)
    return b.getvalue()


A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
ODF = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
       'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"')


def test_pptx_por_slide_com_notas() -> None:
    from app.extract.parsers import parse

    slide = lambda t: f'<p:sld xmlns:p="x" {A}><a:p><a:r><a:t>{t}</a:t></a:r></a:p></p:sld>'  # noqa: E731
    data = _zip({"ppt/slides/slide2.xml": slide("Resultados Q3"), "ppt/slides/slide1.xml": slide("Capa"),
                 "ppt/slides/slide10.xml": slide("Fim"),
                 "ppt/notesSlides/notesSlide2.xml": f'<n {A}><a:t>Falar da meta de 18%</a:t></n>'})
    ex = parse(data, "deck.pptx", "application/octet-stream")
    assert ex.kind == "docx" and ex.pages == 3
    assert [lab for lab, _ in ex.parts] == ["slide 1", "slide 2", "slide 10"]
    assert "Falar da meta de 18%" in ex.parts[1][1]


def test_opendocument_texto_apresentacao_e_planilha() -> None:
    from app.extract.parsers import parse

    odt = _zip({"content.xml": f'<office:document-content {ODF}><office:body><office:text><text:h>Contrato</text:h>'
                               '<text:p>Prazo de 30 dias.</text:p></office:text></office:body></office:document-content>'})
    assert "## Contrato\nPrazo de 30 dias." in parse(odt, "c.odt", "").parts[0][1]
    odp = _zip({"content.xml": f'<r {ODF}><draw:page><text:p>Um</text:p></draw:page><draw:page><text:p>Dois</text:p></draw:page></r>'})
    assert [p[1] for p in parse(odp, "s.odp", "").parts] == ["Um", "Dois"]
    ods = _zip({"content.xml": f'<r {ODF}><table:table table:name="Vendas">'
                               '<table:table-row><table:table-cell><text:p>Região</text:p></table:table-cell><table:table-cell><text:p>Valor</text:p></table:table-cell></table:table-row>'
                               '<table:table-row><table:table-cell><text:p>Sul</text:p></table:table-cell>'
                               '<table:table-cell office:value="1300.5"><text:p>1.300,50</text:p></table:table-cell></table:table-row>'
                               '</table:table></r>'})
    ex = parse(ods, "v.ods", "")
    assert ex.kind == "table" and ex.sheets[0].name == "Vendas" and ex.sheets[0].rows[0][1] == 1300.5


def test_rtf_eml_ipynb_epub_html() -> None:
    import json
    from email.message import EmailMessage

    from app.extract.parsers import parse

    rtf = rb"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0 Relat\'f3rio de \b vendas\b0\par Total: 3.600\par}"
    assert "Relatório de vendas\nTotal: 3.600" in parse(rtf, "r.rtf", "").parts[0][1]

    m = EmailMessage()
    m["From"], m["Subject"] = "ana@x.com", "Pedido 778"
    m.set_content("Segue o pedido.")
    m.add_attachment(b"x", maintype="application", subtype="pdf", filename="pedido.pdf")
    t = parse(m.as_bytes(), "msg.eml", "").parts[0][1]
    assert "Subject: Pedido 778" in t and "Segue o pedido." in t and "pedido.pdf" in t

    nb = {"cells": [{"cell_type": "markdown", "source": ["# Análise"]}, {"cell_type": "code", "source": "x = 1", "outputs": [{"text": "1"}]}]}
    t = parse(json.dumps(nb).encode(), "a.ipynb", "").parts[0][1]
    assert "célula 1 (markdown)" in t and "x = 1" in t and "[saída] 1" in t

    ep = _zip({"OEBPS/cap1.xhtml": "<html><body><h1>Capítulo 1</h1><p>Era uma vez.</p></body></html>", "OEBPS/nav.xhtml": "<nav>índice</nav>"})
    ex = parse(ep, "livro.epub", "")
    assert len(ex.parts) == 1 and "Era uma vez." in ex.parts[0][1]

    html = b"<html><head><style>p{color:red}</style><script>alert(1)</script></head><body><h2>Oi</h2><p>Texto &amp; mais</p></body></html>"
    t = parse(html, "p.html", "").parts[0][1]
    assert "Texto & mais" in t and "alert" not in t and "color" not in t


def test_zip_com_documentos_dentro() -> None:
    from app.extract.parsers import parse

    data = _zip({"pasta/nota.txt": "Prazo de entrega 15 dias", "pasta/tabela.csv": "a;b\n1;2\n", "../fora.txt": "x",
                 "__MACOSX/._nota.txt": "lixo", "pasta/foto.bin": b"\x00\x01"})
    ex = parse(data, "docs.zip", "application/zip")
    labels = " ".join(lab for lab, _ in ex.parts)
    assert "pasta/nota.txt" in labels and "pasta/tabela.csv" in labels and "fora" not in labels and "MACOSX" not in labels
    assert ex.note and "1 arquivo(s) do zip não lidos" in ex.note


def test_formatos_antigos_avisam() -> None:
    from app.extract.parsers import parse

    assert "pptx" in (parse(b"x", "a.ppt", "").note or "") and "JPG" in (parse(b"x", "f.heic", "").note or "")


def test_previa_formatada_markdown_word_e_planilha() -> None:
    import io

    from docx import Document

    from app.render import preview

    md = b"# T\n\n| ID | Teste |\n|---|---|\n| `G-01` | `a.py::x`<br>`b.py::y` |\n\n<script>alert(1)</script>\n\n[x](javascript:alert(2))"
    h = preview(md, "trace.md")["html"]
    assert "<table>" in h and "<br>" in h and "<code>G-01</code>" in h
    assert "<script>" not in h and "javascript:" not in h and "── seção" not in h

    d = Document()
    d.add_heading("Política", 1)
    d.add_paragraph("Texto.")
    d.add_paragraph("um", style="List Bullet")
    d.add_paragraph("dois", style="List Bullet")
    b = io.BytesIO()
    d.save(b)
    h = preview(b.getvalue(), "p.docx", "dark")["html"]
    assert "<h2>Política</h2>" in h and "<ul><li>um</li><li>dois</li></ul>" in h and "color-scheme:dark" in h

    h = preview(b"a;b\n1;2\n", "t.csv")["html"]
    assert "<th>a</th>" in h and "<td>2</td>" in h
